"""Extract supplied DOCX rows without inferring clinical attributes."""
import re
import docx
import pypinyin
from docx.text.paragraph import Paragraph
from docx.table import Table

def extract_notes(path):
    notes = {}
    part = 'WESTERN'
    for index, paragraph in enumerate(docx.Document(path).paragraphs, 1):
        text = clean_text(paragraph.text)
        if text.replace(' ', '') == '（二）中成药':
            part = 'CHINESE_PATENT'
        match = re.match(r'注释\s*(\d+)\s*[：:]', text)
        if match:
            notes.setdefault((part, int(match.group(1))), {'text':text, 'location':f'paragraph:{index}'})
    return notes

def clean_text(text):
    if not text:
        return ""
    text = text.replace('\ufb01', 'fi').replace('\ufb02', 'fl').replace('\ufb00', 'ff')
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    text = re.sub(r'[ \t\u3000\u00a0]+', ' ', text)
    return text.strip()

def clean_cat(text):
    if not text:
        return ""
    # Strip intra-character spaces like '（一）青 霉 素 类' -> '（一）青霉素类'
    # but keep the structure
    clean = clean_text(text)
    return re.sub(r'\s+', '', clean)

def clean_spec_lines(text):
    if not text:
        return []
    lines = [l.strip() for l in text.split('\n') if l.strip()]
    merged = []
    for l in lines:
        if not merged:
            merged.append(l)
        elif ('：' in l or ':' in l) and not l.startswith('（相当') and not l.startswith('(相当'):
            merged.append(l)
        else:
            prev = merged[-1]
            if prev.endswith('、') or prev.endswith('，') or prev.endswith(','):
                merged[-1] = prev + ' ' + l
            else:
                merged[-1] = prev + l
    return merged

def clean_spec_display(text):
    merged = clean_spec_lines(text)
    return '； '.join(merged)

def get_pinyin_code(text):
    if not text:
        return ""
    clean_name = re.sub(r'[^\u4e00-\u9fa5A-Za-z0-9]', '', text)
    if not clean_name:
        return ""
    pinyin_list = pypinyin.pinyin(clean_name, style=pypinyin.Style.FIRST_LETTER)
    code = ''.join([item[0].upper() for item in pinyin_list if item])
    return code[:32]

# Specific known edge cases for western names
WESTERN_SPECIAL_CASES = {
    '35': ('结核病用药', ''),
    '51': ('艾滋病用药', ''),
    '56': ('青蒿素类药物', ''),
    '84': ('美沙拉秦（美沙拉嗪）', 'Mesalazine'),
    '155': ('地尔硫䓬', 'Diltiazem'),
    '207': ('桉柠蒎', 'Eucalyptol, Limonene and Pinene'),
    '225': ('布地格福', 'Budesonide, Glycopyrronium and Formoterol'),
    '256': ('地衣芽孢杆菌活菌', 'Bacillus Licheniformis, Live'),
    '259': ('双歧杆菌四联活菌', 'Combined Bifidobacterium, Lactobacillus, Enterococcus and Bacillus Cereus, Live'),
    '269': ('坦洛新（坦索罗辛）', 'Tamsulosin'),
    '292': ('血友病用药', ''),
    '300': ('羟乙基淀粉 130/0.4', 'Hydroxyethyl Starch 130/0.4'),
    '384': ('三氧化二砷（亚砷酸）', 'Arsenic Trioxide (Arsenious Acid)'),
    '403': ('多种维生素（12）', 'Multivitamin (12)'),
    '406': ('复方氨基酸 18AA', 'Compound Amino Acid 18AA'),
    '407': ('脂肪乳氨基酸葡萄糖', 'Fat Emulsion, Amino Acids and Glucose'),
    '409': ('整蛋白型肠内营养剂（粉剂）', 'Intact Protein Enteral Nutrition Powder'),
    '416': ('乳酸钠林格', "Sodium Lactate Ringer's"),
    '429': ('抗蛇毒血清', ''),
    '431': ('国家免疫规划用疫苗', ''),
    '471': ('避孕药', ''),
}

def split_western_name(seq_str, raw_name):
    clean_raw = clean_text(raw_name)
    vitamin = re.fullmatch(r'维生素\s*([A-Z]\d*)\s+(Vitamin\s+[A-Z]\d*)', clean_raw)
    if vitamin:
        return '维生素' + vitamin.group(1), vitamin.group(2)
    num_str = re.sub(r'[^\d]', '', seq_str)
    if num_str in WESTERN_SPECIAL_CASES:
        return WESTERN_SPECIAL_CASES[num_str]

    # Check if there is an English INN transition
    cn_chars = list(re.finditer(r'[\u4e00-\u9fff\u3400-\u4dbf]', clean_raw))
    if not cn_chars:
        return clean_raw, ''

    last_cn_idx = cn_chars[-1].end()
    cn_part = clean_raw[:last_cn_idx].strip()
    rest = clean_raw[last_cn_idx:].strip()

    # Letter/number suffixes belong to the Chinese drug name when followed by an INN.
    # Source examples: 两性霉素 B\nAmphotericin B and 碳酸钙 D3\nCalcium Carbonate...
    suffix = re.fullmatch(r'([A-Z]\d*)\s+([A-Za-z][\s\S]*)', rest)
    if suffix:
        cn_part += suffix.group(1)
        rest = suffix.group(2)

    # Handle closing brackets after Chinese part: e.g. （嗪） Mesalazine
    if rest.startswith('）') or rest.startswith(')'):
        cn_part += rest[0]
        rest = rest[1:].strip()
    if rest.startswith('］') or rest.startswith(']'):
        cn_part += rest[0]
        rest = rest[1:].strip()

    return cn_part, rest

def extract_rows(path):
    doc = docx.Document(path)
    body = doc.element.body
    current_part = None
    current_major_cat = ''
    current_sub_cat = ''
    current_tcm_func = ''

    western_items = []
    tcm_items = []

    table_index = 0
    for child in body:
        tag = child.tag.split('}')[-1]
        if tag == 'tbl':
            table_index += 1
        if tag == 'p':
            p = Paragraph(child, doc)
            t = p.text.replace(' ', '').strip()
            if not t:
                continue
            if '第一部分' in t or '化学药品和生物制品' in t:
                current_part = 'WESTERN'
                current_major_cat = ''
                current_sub_cat = ''
            elif '第二部分' in t or '中成药' in t:
                current_part = 'CHINESE_PATENT'
                current_major_cat = ''
                current_sub_cat = ''
                current_tcm_func = ''
            elif '第三部分' in t or '中药饮片' in t:
                current_part = 'HERBAL'
                current_major_cat = ''
                current_sub_cat = ''
            elif re.match(r'^[一二三四五六七八九十]+、', t):
                current_major_cat = clean_cat(p.text)
                current_sub_cat = ''
                current_tcm_func = ''
        elif tag == 'tbl' and current_part in ['WESTERN', 'CHINESE_PATENT']:
            tbl = Table(child, doc)
            for row_index, row in enumerate(tbl.rows, 1):
                raw_cells = [clean_text(c.text) for c in row.cells]
                if not any(raw_cells):
                    continue
                first_raw = raw_cells[0].replace(' ', '').strip()
                if first_raw in ['序号', '编号']:
                    continue

                # Subcategory header check
                if re.match(r'^[（(][一二三四五六七八九十\d]+[）)]', first_raw):
                    current_sub_cat = clean_cat(raw_cells[0])
                    current_tcm_func = ''
                    continue

                # Sequence row
                if re.search(r'\d+', first_raw):
                    seq_val = clean_text(raw_cells[0])
                    num_only = int(re.sub(r'[^\d]', '', seq_val))
                    is_duplicate = '*' in seq_val

                    if current_part == 'WESTERN':
                        raw_name = raw_cells[1] if len(raw_cells) > 1 else ''
                        spec = raw_cells[2] if len(raw_cells) > 2 else ''
                        note = raw_cells[3] if len(raw_cells) > 3 else ''
                        has_guidance = '△' in seq_val or '△' in note or '△' in spec

                        cn_name, en_name = split_western_name(seq_val, raw_name)
                        pinyin_code = get_pinyin_code(cn_name)
                        code = f"MED-2026-W{num_only:03d}"

                        western_items.append({
                            'sourceLocation': f'table:{table_index}/row:{row_index}',
                            'rawCells': raw_cells,
                            'part': 'WESTERN',
                            'type_name': '西药',
                            'sd_med_type': 'WESTERN',
                            'major_cat': current_major_cat,
                            'sub_cat': current_sub_cat,
                            'func': '',
                            'seq': seq_val,
                            'num': num_only,
                            'is_duplicate': is_duplicate,
                            'has_guidance': has_guidance,
                            'code': code,
                            'cn_name': cn_name,
                            'en_name': en_name,
                            'pinyin_code': pinyin_code,
                            'spec': spec,
                            'note': note
                        })
                    elif current_part == 'CHINESE_PATENT':
                        func = clean_text(raw_cells[1]) if len(raw_cells) > 1 else ''
                        if func:
                            current_tcm_func = func
                        else:
                            func = current_tcm_func

                        raw_name = raw_cells[2] if len(raw_cells) > 2 else ''
                        spec = raw_cells[3] if len(raw_cells) > 3 else ''
                        note = raw_cells[4] if len(raw_cells) > 4 else ''
                        has_guidance = '△' in seq_val or '△' in note or '△' in spec

                        cn_name = raw_name
                        pinyin_code = get_pinyin_code(cn_name)
                        code = f"MED-2026-T{num_only:03d}"

                        tcm_items.append({
                            'sourceLocation': f'table:{table_index}/row:{row_index}',
                            'rawCells': raw_cells,
                            'part': 'CHINESE_PATENT',
                            'type_name': '中成药',
                            'sd_med_type': 'CHINESE_PATENT',
                            'major_cat': current_major_cat,
                            'sub_cat': current_sub_cat,
                            'func': func,
                            'seq': seq_val,
                            'num': num_only,
                            'is_duplicate': is_duplicate,
                            'has_guidance': has_guidance,
                            'code': code,
                            'cn_name': cn_name,
                            'en_name': '',
                            'pinyin_code': pinyin_code,
                            'spec': spec,
                            'note': note
                        })

    return western_items + tcm_items
