#!/usr/bin/env python3
"""
Build primary care medical service catalog Flyway migrations and seed SQL.
Source: docs/基层医疗收费项目清单_综合医疗服务类.xlsx
Generates:
  - backend/src/main/resources/db/migration/V1_80_0__primary_care_medical_services.sql (PostgreSQL / H2)
  - backend/src/main/resources/db/oracle/V1_80_0__primary_care_medical_services.sql (Oracle)
"""

import os
import sys
import openpyxl
import pypinyin
from decimal import Decimal

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
EXCEL_PATH = os.path.join(ROOT, "docs", "基层医疗收费项目清单_综合医疗服务类.xlsx")
PG_SQL_PATH = os.path.join(ROOT, "backend", "src", "main", "resources", "db", "migration", "V1_80_0__primary_care_medical_services.sql")
ORA_SQL_PATH = os.path.join(ROOT, "backend", "src", "main", "resources", "db", "oracle", "V1_80_0__primary_care_medical_services.sql")

TENANT_ID = 362387869790209
ORG_ID = 362387869790211
DEPT_GP_ID = 362387869790212       # 全科门诊
DEPT_WARD_ID = 362387869898501     # 综合病区
USER_ID = 362387869790222          # 系统管理人员
BASE_ID = 362387890000000

# Item Type IDs
ITEM_TYPE_SERVICE = 362387869797002      # 诊疗服务
ITEM_TYPE_LAB = 362387869797003          # 检验项目
ITEM_TYPE_EXAM = 362387869797004         # 检查项目
ITEM_TYPE_PROCEDURE = 362387869797005    # 处置操作
ITEM_TYPE_TREATMENT = 362387869797006    # 治疗项目

UNIT_MAPPING = {
    "每部位": "部位",
    "每肢": "肢",
    "科/次": "次",
    "次·人": "人次",
    "天": "日",
    "": "次",
    None: "次",
}

TREATMENT_CATEGORIES = {
    "抢救费", "氧气吸入", "注射", "清创（缝合）", "换药（包括拆线、术后清创换药）",
    "雾化吸入", "鼻饲管置管", "胃肠减压", "洗胃", "物理降温", "坐浴", "冷热湿敷",
    "引流管冲洗", "灌肠", "导尿", "肛管排气", "院前急救费", "救护车费"
}

def clean_unit(unit):
    if unit in UNIT_MAPPING:
        return UNIT_MAPPING[unit]
    u = (unit or "").strip()
    return UNIT_MAPPING.get(u, u if u else "次")

def get_pinyin_initials(text):
    if not text:
        return ""
    letters = pypinyin.pinyin(text, style=pypinyin.Style.FIRST_LETTER)
    return "".join([item[0].upper() for item in letters if item[0].isalnum()])

def sql_escape(value):
    if value is None:
        return "null"
    s = str(value).replace("'", "''")
    return f"'{s}'"

def parse_items():
    wb = openpyxl.load_workbook(EXCEL_PATH)
    items = []
    seen_codes = set()

    # Sheet 1: 综合医疗服务类清单
    s1 = wb["综合医疗服务类清单"]
    for r in list(s1.iter_rows(values_only=True))[1:]:
        if not r or not r[1]:
            continue
        code = str(r[1]).strip()
        if code in seen_codes:
            continue
        seen_codes.add(code)
        name = str(r[2]).strip()
        unit = clean_unit(r[3])
        raw_price = r[4]
        insurance_type = str(r[5]).strip() if r[5] else None
        category = str(r[6]).strip() if r[6] else ""
        grassroots = str(r[7]).strip() if r[7] else None
        notes = str(r[8]).strip() if r[8] else None
        connotation = str(r[9]).strip() if len(r) > 9 and r[9] else None

        items.append({
            "code": code,
            "name": name,
            "unit": unit,
            "price": raw_price,
            "insurance_type": insurance_type,
            "category": category,
            "grassroots": grassroots,
            "notes": notes,
            "connotation": connotation,
            "sheet": 1
        })

    # Sheet 2: 常见检验检查项目
    s2 = wb["常见检验检查项目"]
    for r in list(s2.iter_rows(values_only=True))[1:]:
        if not r or not r[1]:
            continue
        code = str(r[1]).strip()
        if code in seen_codes:
            continue
        seen_codes.add(code)
        name = str(r[2]).strip()
        unit = clean_unit(r[3])
        raw_price = r[4]
        insurance_type = str(r[5]).strip() if r[5] else None
        category = str(r[6]).strip() if r[6] else ""
        connotation = str(r[7]).strip() if len(r) > 7 and r[7] else None
        notes = str(r[8]).strip() if len(r) > 8 and r[8] else None

        items.append({
            "code": code,
            "name": name,
            "unit": unit,
            "price": raw_price,
            "insurance_type": insurance_type,
            "category": category,
            "grassroots": None,
            "notes": notes,
            "connotation": connotation,
            "sheet": 2
        })

    return items

def derive_metadata(item):
    code = item["code"]
    cat = item["category"]

    # Default values
    price_unit = Decimal("0.000000")
    price_doc = "ZJ-GUIDELINE-PRICE"
    price_reason = "政府指导价/分级定价待机构核定"

    if item["price"] is not None:
        try:
            price_unit = Decimal(str(item["price"])).quantize(Decimal("0.000000"))
            price_doc = "ZJ-YB-2024"
            price_reason = "浙江省2024基准价格"
        except Exception:
            pass

    # Classify by code and category
    if code.startswith(("21", "22", "23", "24")):
        svc_type = "EXAMINATION"
        item_type_id = ITEM_TYPE_EXAM
        acctg_cat = "EXAMINATION"
        is_medical_tech = 1
        usage_type = "COMMON"
        dept_id = DEPT_GP_ID
    elif code.startswith(("25", "26")):
        svc_type = "LABORATORY"
        item_type_id = ITEM_TYPE_LAB
        acctg_cat = "LABORATORY"
        is_medical_tech = 1
        usage_type = "COMMON"
        dept_id = DEPT_GP_ID
    elif code.startswith("1201") or cat == "护理费":
        svc_type = "NURSING"
        item_type_id = ITEM_TYPE_SERVICE
        acctg_cat = "NURSING"
        is_medical_tech = 0
        usage_type = "INPATIENT"
        dept_id = DEPT_WARD_ID
    elif code.startswith(("1101", "1102")) or cat in ["挂号费", "诊查费", "一般诊疗费"]:
        svc_type = "PROCEDURE"
        item_type_id = ITEM_TYPE_PROCEDURE
        acctg_cat = "REGISTRATION"
        is_medical_tech = 0
        usage_type = "OUTPATIENT"
        dept_id = DEPT_GP_ID
    elif code.startswith("1109") or cat in ["床位费", "中心监护病房", "空调费"]:
        svc_type = "OTHER"
        item_type_id = ITEM_TYPE_SERVICE
        acctg_cat = "BED"
        is_medical_tech = 0
        usage_type = "INPATIENT"
        dept_id = DEPT_WARD_ID
    elif cat in TREATMENT_CATEGORIES or code.startswith(("31", "32", "33", "34", "41", "42", "43", "44", "45", "46", "47")):
        svc_type = "TREATMENT"
        item_type_id = ITEM_TYPE_TREATMENT
        acctg_cat = "TREATMENT"
        is_medical_tech = 0
        usage_type = "COMMON"
        dept_id = DEPT_GP_ID
    else:
        svc_type = "OTHER"
        item_type_id = ITEM_TYPE_SERVICE
        acctg_cat = "OTHER"
        is_medical_tech = 0
        usage_type = "COMMON"
        dept_id = DEPT_GP_ID

    item["price_unit"] = price_unit
    item["price_doc"] = price_doc
    item["price_reason"] = price_reason
    item["svc_type"] = svc_type
    item["item_type_id"] = item_type_id
    item["acctg_cat"] = acctg_cat
    item["is_medical_tech"] = is_medical_tech
    item["usage_type"] = usage_type
    item["dept_id"] = dept_id
    item["pinyin"] = get_pinyin_initials(item["name"])
    return item

def generate_sql(items, dialect="oracle"):
    is_oracle = (dialect == "oracle")
    lines = [
        "-- V1_80_0__primary_care_medical_services.sql",
        f"-- Ingest 409 standard primary care medical service catalog items for tenant {TENANT_ID}",
        f"-- Generated automatically from {os.path.basename(EXCEL_PATH)}",
        "",
    ]

    for idx, item in enumerate(items, start=1):
        cat_item_id = BASE_ID + idx
        price_id = BASE_ID + 1000000 + idx
        org_item_id = BASE_ID + 2000000 + idx
        search_tnt_id = BASE_ID + 3000000 + idx
        search_org_id = BASE_ID + 4000000 + idx
        attr_subject_id = BASE_ID + 5000000 + idx

        code = sql_escape(item["code"])
        name = sql_escape(item["name"])
        unit = sql_escape(item["unit"])
        category = sql_escape(item["category"])
        pinyin = sql_escape(item["pinyin"])
        attention = sql_escape(item["notes"])
        exam_notes = sql_escape(item["connotation"])

        # boolean formatting
        b_true = "1" if is_oracle else "true"
        b_false = "0" if is_oracle else "false"
        b_med_tech = str(item["is_medical_tech"]) if is_oracle else ("true" if item["is_medical_tech"] else "false")

        # timestamp formatting
        if is_oracle:
            ts_created = "to_timestamp_tz('2024-01-01 00:00:00 +08:00', 'YYYY-MM-DD HH24:MI:SS TZH:TZM')"
            from_clause = "from dual where not exists"
        else:
            ts_created = "timestamp with time zone '2024-01-01 00:00:00+08'"
            from_clause = "where not exists"

        # 1. RHN_BD_CATALOG_ITEM
        lines.append(
            f"insert into RHN_BD_CATALOG_ITEM select {cat_item_id}, 0, {TENANT_ID}, {code}, {name}, 'SERVICE', {unit}, "
            f"{b_true}, {b_true}, {b_false}, 'ACTIVE', date '2024-01-01', null, {ts_created}, {USER_ID}, {ts_created}, {USER_ID}, "
            f"{item['item_type_id']}, null {from_clause} (select 1 from RHN_BD_CATALOG_ITEM where ID_CATALOG_ITEM = {cat_item_id});"
        )

        # 2. RHN_BD_SVC_ITEM
        lines.append(
            f"insert into RHN_BD_SVC_ITEM select {cat_item_id}, {TENANT_ID}, '{item['svc_type']}', {category}, '{item['usage_type']}', "
            f"{b_med_tech}, {b_false}, {b_true}, null, null, '{item['acctg_cat']}', null, null, null, null, null, "
            f"{b_false}, {attention}, {exam_notes} {from_clause} (select 1 from RHN_BD_SVC_ITEM where ID_CATALOG_ITEM = {cat_item_id});"
        )

        # 3. RHN_BD_CATALOG_PRICE (Provincial baseline price, organizationId = null for tenant baseline)
        lines.append(
            f"insert into RHN_BD_CATALOG_PRICE select {price_id}, 0, {TENANT_ID}, {cat_item_id}, null, null, 'SALE', "
            f"{item['price_unit']:.6f}, 'CNY', '{item['price_doc']}', '{item['price_reason']}', date '2024-01-01', null, "
            f"'ACTIVE', {ts_created}, {USER_ID}, {ts_created}, {USER_ID}, null {from_clause} "
            f"(select 1 from RHN_BD_CATALOG_PRICE where ID_CATALOG_PRICE = {price_id});"
        )

        # 4. RHN_BD_ORG_CATALOG_ITEM (Adopted for default health center 362387869790211)
        lines.append(
            f"insert into RHN_BD_ORG_CATALOG_ITEM select {org_item_id}, 0, {TENANT_ID}, {ORG_ID}, {cat_item_id}, {item['dept_id']}, "
            f"{code}, {name}, {b_true}, {b_true}, {b_true}, {b_false}, {b_false}, {b_false}, {b_false}, 'ACTIVE', "
            f"date '2024-01-01', null, {ts_created}, {USER_ID}, {ts_created}, {USER_ID}, null {from_clause} "
            f"(select 1 from RHN_BD_ORG_CATALOG_ITEM where ID_ORG_CATALOG_ITEM = {org_item_id});"
        )

        # 5. RHN_BD_ITEM_ATTR_SUBJECT
        subject_urn = sql_escape(f"TENANT:{TENANT_ID}/CATALOG_ITEM:{cat_item_id}")
        lines.append(
            f"insert into RHN_BD_ITEM_ATTR_SUBJECT select {attr_subject_id}, {TENANT_ID}, 'CATALOG_ITEM', {subject_urn}, "
            f"null, null, {cat_item_id}, null, {ts_created}, {USER_ID} {from_clause} "
            f"(select 1 from RHN_BD_ITEM_ATTR_SUBJECT where ID_ITEM_ATTR_SUBJECT = {attr_subject_id});"
        )

        # 6. RHN_BD_SEARCH_ENTRY (Tenant Level)
        lines.append(
            f"insert into RHN_BD_SEARCH_ENTRY select {search_tnt_id}, 0, 'TENANT', {TENANT_ID}, {TENANT_ID}, 'CATALOG_ITEM', "
            f"{cat_item_id}, 'CANONICAL', 'service:canonical', {name}, {pinyin}, null, null, {b_true}, 'search-code-v1', "
            f"'ACTIVE', {ts_created}, {USER_ID}, {ts_created}, {USER_ID} {from_clause} "
            f"(select 1 from RHN_BD_SEARCH_ENTRY where ID_SEARCH_ENTRY = {search_tnt_id});"
        )

        # 7. RHN_BD_SEARCH_ENTRY (Organization Level)
        lines.append(
            f"insert into RHN_BD_SEARCH_ENTRY select {search_org_id}, 0, 'ORGANIZATION', {ORG_ID}, {TENANT_ID}, 'CATALOG_ITEM', "
            f"{cat_item_id}, 'LOCAL_NAME', 'adoption:local', {name}, {pinyin}, null, null, {b_true}, 'search-code-v1', "
            f"'ACTIVE', {ts_created}, {USER_ID}, {ts_created}, {USER_ID} {from_clause} "
            f"(select 1 from RHN_BD_SEARCH_ENTRY where ID_SEARCH_ENTRY = {search_org_id});"
        )

    return "\n".join(lines) + "\n"

def main():
    print(f"Loading items from {EXCEL_PATH}...")
    items = parse_items()
    print(f"Total unique items parsed: {len(items)}")

    for item in items:
        derive_metadata(item)

    # Summary
    by_type = {}
    for item in items:
        t = item["svc_type"]
        by_type[t] = by_type.get(t, 0) + 1
    print("Items by Service Type:", by_type)

    os.makedirs(os.path.dirname(PG_SQL_PATH), exist_ok=True)
    os.makedirs(os.path.dirname(ORA_SQL_PATH), exist_ok=True)

    print(f"Generating PostgreSQL/H2 migration -> {PG_SQL_PATH}...")
    pg_sql = generate_sql(items, dialect="postgresql")
    with open(PG_SQL_PATH, "w", encoding="utf-8") as f:
        f.write(pg_sql)

    print(f"Generating Oracle migration -> {ORA_SQL_PATH}...")
    ora_sql = generate_sql(items, dialect="oracle")
    with open(ORA_SQL_PATH, "w", encoding="utf-8") as f:
        f.write(ora_sql)

    print(f"Migration scripts generated successfully! (Lines: {len(ora_sql.splitlines())})")

if __name__ == "__main__":
    main()
